"""One-off generator: Rosary Sequence .docx from src/data/prayers.json + joyful mysteries."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
PRAYERS = json.loads((ROOT / "src/data/prayers.json").read_text(encoding="utf-8"))
JOYFUL = json.loads((ROOT / "src/data/mystery-groups/joyful.json").read_text(encoding="utf-8"))
ROSARY = json.loads((ROOT / "src/data/rosary-types/rosary-joyful.json").read_text(encoding="utf-8"))


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_note(doc, text, italic=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(11)
    return p


def add_prayer_block(doc, prayer_id, repeat_note=None):
    p = PRAYERS[prayer_id]
    add_heading(doc, p["title"], level=2)
    doc.add_paragraph(p["text"])
    if repeat_note:
        add_note(doc, repeat_note)


def main():
    doc = Document()
    t = doc.add_heading("Rosary Sequence", 0)
    t.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph(
        "This document lists every prayer said in one complete Holy Rosary (Joyful Mysteries), "
        "in order. Where the Hail Mary appears, the note indicates how many times to repeat "
        "that same prayer — it is not a different prayer each time."
    )

    # Opening (crucifix through Fatima, then decades)
    add_heading(doc, "Opening", level=1)
    for pid in ROSARY["openingPrayerIds"]:
        add_prayer_block(doc, pid)

    # Five decades
    add_heading(doc, "Five decades", level=1)
    for i, m in enumerate(JOYFUL["mysteries"], start=1):
        add_heading(doc, f"Decade {i}", level=2)
        add_prayer_block(doc, m["announcementPrayerId"])

        ids = m["decadePrayerIds"]
        if ids[0] != "our-father":
            raise SystemExit(f"Unexpected decade layout: {m['id']}")
        add_prayer_block(doc, "our-father")
        add_heading(doc, PRAYERS["hail-mary"]["title"], level=2)
        doc.add_paragraph(PRAYERS["hail-mary"]["text"])
        add_note(
            doc,
            "Repeat 10×: say this same Hail Mary ten times while meditating on the "
            "mystery above. Each repetition is the same prayer, not ten different prayers.",
        )
        tail = [pid for pid in ids[1:] if pid != "hail-mary"]
        for pid in tail:
            add_prayer_block(doc, pid)

    # Closing
    add_heading(doc, "Closing", level=1)
    for pid in ROSARY["closingPrayerIds"]:
        add_prayer_block(doc, pid)

    out = ROOT / "Rosary-Sequence.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
