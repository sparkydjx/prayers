"""Generate Rosary-Mysteries.docx: sets, weekday schedule, decade ↔ mystery numbering."""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
PRAYERS = json.loads((ROOT / "src/data/prayers.json").read_text(encoding="utf-8"))

# USCCB-style default assignment (ordinary time). Some communities vary in Advent/Lent.
WEEKDAY_ROWS = [
    ("Sunday", "Glorious Mysteries"),
    ("Monday", "Joyful Mysteries"),
    ("Tuesday", "Sorrowful Mysteries"),
    ("Wednesday", "Glorious Mysteries"),
    ("Thursday", "Luminous Mysteries"),
    ("Friday", "Sorrowful Mysteries"),
    ("Saturday", "Joyful Mysteries"),
]

MYSTERY_SETS = [
    {
        "title": "Joyful Mysteries",
        "days": "Monday and Saturday",
        "prayer_keys": [
            "mystery-joyful-1",
            "mystery-joyful-2",
            "mystery-joyful-3",
            "mystery-joyful-4",
            "mystery-joyful-5",
        ],
    },
    {
        "title": "Luminous Mysteries",
        "days": "Thursday",
        "items": [
            (
                "First Luminous Mystery — The Baptism of the Lord",
                "Christ is baptized by John in the Jordan; the Father proclaims Him beloved Son.",
            ),
            (
                "Second Luminous Mystery — The Wedding at Cana",
                "Jesus performs His first sign at Mary’s request, manifesting His glory.",
            ),
            (
                "Third Luminous Mystery — The Proclamation of the Kingdom",
                "Jesus calls to conversion and announces the good news of the Kingdom.",
            ),
            (
                "Fourth Luminous Mystery — The Transfiguration",
                "Jesus is revealed in glory to Peter, James, and John on the holy mountain.",
            ),
            (
                "Fifth Luminous Mystery — The Institution of the Eucharist",
                "Jesus offers His Body and Blood at the Last Supper as the memorial of His sacrifice.",
            ),
        ],
    },
    {
        "title": "Sorrowful Mysteries",
        "days": "Tuesday and Friday",
        "items": [
            (
                "First Sorrowful Mystery — The Agony in the Garden",
                "Jesus prays in Gethsemane, accepting the Father’s will in His suffering.",
            ),
            (
                "Second Sorrowful Mystery — The Scourging at the Pillar",
                "Jesus is cruelly whipped at the order of Pilate.",
            ),
            (
                "Third Sorrowful Mystery — The Crowning with Thorns",
                "Jesus is mocked as King with a crown of thorns.",
            ),
            (
                "Fourth Sorrowful Mystery — The Carrying of the Cross",
                "Jesus carries the Cross to Calvary for our salvation.",
            ),
            (
                "Fifth Sorrowful Mystery — The Crucifixion",
                "Jesus dies on the Cross, offering His life for the world.",
            ),
        ],
    },
    {
        "title": "Glorious Mysteries",
        "days": "Sunday and Wednesday",
        "items": [
            (
                "First Glorious Mystery — The Resurrection",
                "Christ rises triumphant from the dead on the third day.",
            ),
            (
                "Second Glorious Mystery — The Ascension",
                "Jesus ascends into heaven and is seated at the right hand of the Father.",
            ),
            (
                "Third Glorious Mystery — The Descent of the Holy Spirit",
                "The Holy Spirit comes upon Mary and the Apostles at Pentecost.",
            ),
            (
                "Fourth Glorious Mystery — The Assumption",
                "The Blessed Virgin Mary is taken body and soul into heavenly glory.",
            ),
            (
                "Fifth Glorious Mystery — The Coronation of Mary",
                "Mary is crowned Queen of heaven and earth as Mother of the King of kings.",
            ),
        ],
    },
]


def add_schedule_table(doc):
    doc.add_heading("Which mysteries on which days", level=1)
    t = doc.add_table(rows=1 + len(WEEKDAY_ROWS), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    hdr[0].text = "Day"
    hdr[1].text = "Mystery set"
    for i, (day, mset) in enumerate(WEEKDAY_ROWS, start=1):
        row = t.rows[i].cells
        row[0].text = day
        row[1].text = mset
    doc.add_paragraph()


def resolve_mysteries(ms):
    if "prayer_keys" in ms:
        out = []
        for k in ms["prayer_keys"]:
            p = PRAYERS[k]
            out.append((p["title"], p["text"]))
        return out
    return [(t, d) for t, d in ms["items"]]


def main():
    doc = Document()
    h = doc.add_heading("Rosary Mysteries", 0)
    h.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph(
        "The full rosary has twenty mysteries, grouped into four sets of five. When you pray "
        "one rosary, you meditate on one set—the five mysteries assigned to that day (see table "
        "below). In the Rosary Sequence, the decades are numbered in order: after the opening "
        "prayers, the first decade is always decade 1, the second is decade 2, and so on through "
        "decade 5. Each decade matches one mystery of the day’s set: decade 1 goes with mystery 1, "
        "decade 2 with mystery 2, etc. The announcement before each decade names that mystery."
    )

    p = doc.add_paragraph()
    r = p.add_run(
        "Note: In Advent and Lent, some communities pray the Joyful or Sorrowful mysteries on "
        "Sundays instead of the Glorious set. Follow your parish or bishop’s guidance when it differs."
    )
    r.italic = True
    r.font.size = Pt(11)

    doc.add_paragraph()
    add_schedule_table(doc)

    doc.add_heading("Structure of each mystery (in sequence)", level=1)
    doc.add_paragraph(
        "For each mystery number N (1 through 5), the Rosary Sequence contains, in order: "
        "the announcement for mystery N (title + short meditation), one Our Father, ten Hail Marys "
        "(the same prayer repeated), the Glory Be, and the Fatima prayer. Then the next mystery "
        "begins with its announcement, until all five decades are complete, followed by the closing "
        "prayers."
    )

    for ms in MYSTERY_SETS:
        doc.add_heading(ms["title"], level=1)
        p = doc.add_paragraph()
        r = p.add_run(f"Traditionally prayed on: {ms['days']}.")
        r.bold = True

        mysteries = resolve_mysteries(ms)
        for i, (title, text) in enumerate(mysteries, start=1):
            doc.add_heading(f"Mystery {i} — Decade {i} in the Rosary Sequence", level=2)
            doc.add_paragraph(title)
            doc.add_paragraph(text)

    out = ROOT / "Rosary-Mysteries.docx"
    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
